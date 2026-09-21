from __future__ import annotations

import io
from datetime import date, datetime, time, timedelta
from urllib.parse import quote
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import HTMLResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.orm import Session

from .db import get_db
from .models import FuelEvent

router = APIRouter()
TZ = ZoneInfo("Europe/Istanbul")
MAX_REPORT_DAYS = 366

SOURCE_ORDER = ("turpak", "shell", "petrol")
SOURCE_LABELS = {"turpak": "Turpak", "shell": "Shell", "petrol": "Petrol"}
SOURCE_DB_VALUES = {
    "turpak": ("turpak",),
    "shell": ("shell_excel",),
    "petrol": ("petrol",),
}
TURPAK_STATION_MAP = {
    "000003": "90VBX",
    "000004": "91VBX",
    "100007": "33AHZ660",
    "100010": "33AHZ941",
    "100009": "34AHZ947",
    "000005": "AMBAR",
}

EXCEL_HEADERS = ("Станция / танкер", "Госномер", "Время заправки", "Объем, л")
SCREEN_HEADERS = ("Система",) + EXCEL_HEADERS
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _parse_day(value: str) -> date:
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Некорректная дата") from exc


def _selected_sources(raw: str) -> list[str]:
    requested = {item.strip().lower() for item in raw.split(",") if item.strip()}
    selected = [item for item in SOURCE_ORDER if item in requested]
    if not selected:
        raise HTTPException(status_code=400, detail="Выберите хотя бы одну систему")
    return selected

def _station_display(source_key: str, station_code, station_name) -> str:
    name = str(station_name or "").strip()
    code = str(station_code or "").strip()
    if source_key == "turpak" and code:
        return TURPAK_STATION_MAP.get(code, name or code)
    if name:
        return name
    return code or "—"


def _source_key(db_source: str) -> str:
    for source_key, values in SOURCE_DB_VALUES.items():
        if db_source in values:
            return source_key
    return db_source


def _load_rows(db: Session, start_day: date, end_day: date, selected: list[str]) -> list[dict]:
    if end_day < start_day:
        raise HTTPException(status_code=400, detail="Дата окончания раньше даты начала")
    if (end_day - start_day).days + 1 > MAX_REPORT_DAYS:
        raise HTTPException(status_code=400, detail="Период отчета не должен превышать 366 дней")

    start_dt = datetime.combine(start_day, time.min, tzinfo=TZ)
    end_dt = datetime.combine(end_day + timedelta(days=1), time.min, tzinfo=TZ)
    rows: list[dict] = []

    if "turpak" in selected:
        full_turpak = db.execute(
            text(
                """
                SELECT event_dt, plate, liters, station_code, station_name
                FROM public.turpak_fuel_events_all
                WHERE event_dt >= :start_dt AND event_dt < :end_dt
                ORDER BY event_dt, plate, id
                """
            ),
            {"start_dt": start_dt, "end_dt": end_dt},
        ).mappings()
        for event in full_turpak:
            event_dt = event["event_dt"]
            if event_dt is not None and event_dt.tzinfo is not None:
                event_dt = event_dt.astimezone(TZ)
            rows.append({
                "source_key": "turpak",
                "source_label": SOURCE_LABELS["turpak"],
                "station": _station_display("turpak", event["station_code"], event["station_name"]),
                "plate": str(event["plate"] or "").strip() or "—",
                "event_dt": event_dt,
                "liters": float(event["liters"] or 0),
            })

    card_sources = [
        value
        for key in selected
        if key != "turpak"
        for value in SOURCE_DB_VALUES[key]
    ]
    if card_sources:
        items = (
            db.query(FuelEvent)
            .filter(
                FuelEvent.event_dt >= start_dt,
                FuelEvent.event_dt < end_dt,
                FuelEvent.source.in_(card_sources),
            )
            .order_by(FuelEvent.event_dt.asc(), FuelEvent.plate.asc(), FuelEvent.id.asc())
            .all()
        )
        for event in items:
            source_key = _source_key(str(getattr(event, "source", "") or ""))
            event_dt = getattr(event, "event_dt", None)
            if event_dt is not None and event_dt.tzinfo is not None:
                event_dt = event_dt.astimezone(TZ)
            rows.append({
                "source_key": source_key,
                "source_label": SOURCE_LABELS.get(source_key, source_key),
                "station": _station_display(
                    source_key,
                    getattr(event, "station_code", None),
                    getattr(event, "station_name", None),
                ),
                "plate": str(getattr(event, "plate", "") or "").strip() or "—",
                "event_dt": event_dt,
                "liters": float(getattr(event, "liters", 0) or 0),
            })

    rows.sort(key=lambda row: (row["event_dt"] or datetime.min.replace(tzinfo=TZ), row["plate"]))
    return rows


def _autosize(sheet) -> None:
    widths = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)) + 2)
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width, 12), 42)


def _workbook(rows: list[dict], selected: list[str]) -> io.BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    for source_key in selected:
        ws = wb.create_sheet(SOURCE_LABELS[source_key])
        ws.append(list(EXCEL_HEADERS))
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in rows:
            if row["source_key"] != source_key:
                continue
            event_dt = row["event_dt"]
            if event_dt is not None and event_dt.tzinfo is not None:
                event_dt = event_dt.replace(tzinfo=None)
            ws.append([row["station"], row["plate"], event_dt, row["liters"]])
        for cell in ws["C"][1:]:
            cell.number_format = "dd.mm.yyyy hh:mm:ss"
        for cell in ws["D"][1:]:
            cell.number_format = "0.00"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        _autosize(ws)
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream

FUEL_REPORT_HTML = r"""<!doctype html>
<html lang="ru">
<head>
  <meta charset="utf-8"/>
  <meta name="viewport" content="width=device-width,initial-scale=1"/>
  <title>Отчет по топливу</title>
  <style>
    :root{font-family:Inter,Segoe UI,Arial,sans-serif;color:#0f172a;background:#f6f7f9}
    body{margin:0}.container{max-width:1500px;margin:0 auto;padding:24px}
    .card{background:#fff;border:1px solid #e2e8f0;border-radius:20px;padding:20px;margin-bottom:16px;box-shadow:0 8px 24px rgba(15,23,42,.06)}
    .header,.actions,.sources{display:flex;gap:12px;align-items:center;flex-wrap:wrap}.header{justify-content:space-between}
    h1{margin:0}.muted{color:#64748b}.btn,button,input{border:1px solid #cbd5e1;background:#fff;color:#0f172a;padding:10px 14px;border-radius:14px;font-size:14px}
    .btn,button{cursor:pointer;text-decoration:none;font-weight:650}.primary{background:#111827;color:#fff;border-color:#111827}
    .filters{display:grid;grid-template-columns:180px 180px 1fr;gap:14px;margin-top:18px;align-items:end}
    label{display:block;font-size:13px;font-weight:650;margin-bottom:6px}.sources label{display:flex;align-items:center;gap:7px;margin:0}
    input[type=checkbox]{width:auto}.status{margin-top:14px;min-height:20px}.error{color:#b42318}.ok{color:#067647}
    .stats{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin-bottom:14px}.stat{border:1px solid #e2e8f0;border-radius:16px;padding:12px}
    .stat-k{font-size:12px;color:#64748b}.stat-v{font-size:20px;font-weight:750;margin-top:4px}
    .table-wrap{overflow:auto;max-height:68vh;border:1px solid #e2e8f0;border-radius:14px}
    table{width:100%;border-collapse:collapse;font-size:13px}th,td{padding:9px 10px;border-bottom:1px solid #e2e8f0;white-space:nowrap;text-align:left}
    th{position:sticky;top:0;background:#f8fafc;color:#475569}th:last-child,td:last-child{text-align:right}.hidden{display:none}
    @media(max-width:900px){.filters,.stats{grid-template-columns:1fr 1fr}}@media(max-width:560px){.filters,.stats{grid-template-columns:1fr}}
  </style>
</head>
<body><div class="container">
  <div class="card">
    <div class="header">
      <div><div class="muted">Топливный мониторинг</div><h1>Отчет по топливу</h1></div>
      <a class="btn" href="/leadership">← Назад</a>
    </div>
    <div class="filters">
      <div><label for="date-from">Дата от</label><input id="date-from" type="date"/></div>
      <div><label for="date-to">Дата до</label><input id="date-to" type="date"/></div>
      <div><label>Системы</label><div class="sources">
        <label><input type="checkbox" name="source" value="turpak" checked/>Turpak</label>
        <label><input type="checkbox" name="source" value="shell" checked/>Shell</label>
        <label><input type="checkbox" name="source" value="petrol" checked/>Petrol</label>
      </div></div>
    </div>
    <div class="actions" style="margin-top:16px">
      <button id="build-btn" class="primary" type="button">Сформировать отчет</button>
      <button id="excel-btn" type="button" disabled>Скачать Excel</button>
    </div>
    <div id="status" class="status muted"></div>
  </div>
  <div id="result" class="card hidden">
    <div id="stats" class="stats"></div>
    <div class="table-wrap"><table><thead><tr id="head"></tr></thead><tbody id="body"></tbody></table></div>
  </div>
</div>
<script>
const fromInput=document.getElementById('date-from');
const toInput=document.getElementById('date-to');
const buildBtn=document.getElementById('build-btn');
const excelBtn=document.getElementById('excel-btn');
const statusBox=document.getElementById('status');
const result=document.getElementById('result');
const stats=document.getElementById('stats');
const head=document.getElementById('head');
const body=document.getElementById('body');
let lastQuery='';

function esc(v){return String(v==null?'':v).replace(/[&<>'"]/g,function(c){return {'&':'&amp;','<':'&lt;','>':'&gt;',"'":'&#39;','"':'&quot;'}[c]})}
function initDates(){
  const now=new Date(); const y=now.getFullYear(); const m=String(now.getMonth()+1).padStart(2,'0'); const d=String(now.getDate()).padStart(2,'0');
  fromInput.value=y+'-'+m+'-01'; toInput.value=y+'-'+m+'-'+d;
}
function selectedSources(){return Array.from(document.querySelectorAll('input[name=source]:checked')).map(function(x){return x.value})}
function buildQuery(){
  const sources=selectedSources();
  if(!sources.length) throw new Error('Выберите хотя бы одну систему.');
  if(!fromInput.value||!toInput.value) throw new Error('Укажите период.');
  return 'date_from='+encodeURIComponent(fromInput.value)+'&date_to='+encodeURIComponent(toInput.value)+'&sources='+encodeURIComponent(sources.join(','));
}
function render(data){
  head.innerHTML=data.columns.map(function(x){return '<th>'+esc(x)+'</th>'}).join('');
  body.innerHTML=data.rows.map(function(row){return '<tr>'+row.map(function(x){return '<td>'+esc(x)+'</td>'}).join('')+'</tr>'}).join('');
  stats.innerHTML=Object.entries(data.summary).map(function(item){return '<div class="stat"><div class="stat-k">'+esc(item[0])+'</div><div class="stat-v">'+esc(item[1])+'</div></div>'}).join('');
  result.classList.remove('hidden');
}
initDates();
fromInput.addEventListener('change',function(){toInput.min=fromInput.value;if(toInput.value<fromInput.value)toInput.value=fromInput.value});
buildBtn.addEventListener('click',async function(){
  buildBtn.disabled=true; excelBtn.disabled=true; result.classList.add('hidden'); statusBox.className='status muted'; statusBox.textContent='Формирование отчета...';
  try{
    lastQuery=buildQuery();
    const response=await fetch('/fuel-report/data?'+lastQuery,{cache:'no-store'});
    const payload=await response.json();
    if(!response.ok)throw new Error(payload.detail||'Ошибка формирования отчета');
    render(payload); excelBtn.disabled=false; statusBox.className='status ok'; statusBox.textContent='Готово. Заправок: '+payload.rows.length+'.';
  }catch(error){statusBox.className='status error';statusBox.textContent=error.message}
  finally{buildBtn.disabled=false}
});
excelBtn.addEventListener('click',function(){if(lastQuery)window.location.href='/fuel-report/export?'+lastQuery});
</script></body></html>"""

@router.get("/fuel-report", response_class=HTMLResponse, include_in_schema=False)
def fuel_report_page() -> str:
    return FUEL_REPORT_HTML


@router.get("/fuel-report/data")
def fuel_report_data(
    date_from: str = Query(...),
    date_to: str = Query(...),
    sources: str = Query("turpak,shell,petrol"),
    db: Session = Depends(get_db),
):
    start_day = _parse_day(date_from)
    end_day = _parse_day(date_to)
    selected = _selected_sources(sources)
    rows = _load_rows(db, start_day, end_day, selected)
    display_rows = [
        [
            row["source_label"],
            row["station"],
            row["plate"],
            row["event_dt"].strftime("%d.%m.%Y %H:%M:%S") if row["event_dt"] else "",
            round(row["liters"], 2),
        ]
        for row in rows
    ]
    period = start_day.strftime("%d.%m.%Y")
    if end_day != start_day:
        period += "–" + end_day.strftime("%d.%m.%Y")
    return {
        "columns": list(SCREEN_HEADERS),
        "rows": display_rows,
        "summary": {
            "Период": period,
            "Системы": ", ".join(SOURCE_LABELS[key] for key in selected),
            "Заправок": len(rows),
            "Объем, л": f"{sum(row['liters'] for row in rows):,.2f}".replace(",", " "),
        },
    }


@router.get("/fuel-report/export")
def fuel_report_export(
    date_from: str = Query(...),
    date_to: str = Query(...),
    sources: str = Query("turpak,shell,petrol"),
    db: Session = Depends(get_db),
):
    start_day = _parse_day(date_from)
    end_day = _parse_day(date_to)
    selected = _selected_sources(sources)
    rows = _load_rows(db, start_day, end_day, selected)
    stream = _workbook(rows, selected)
    filename = f"Отчет_по_топливу_{start_day.isoformat()}_{end_day.isoformat()}.xlsx"
    ascii_name = f"fuel_report_{start_day.isoformat()}_{end_day.isoformat()}.xlsx"
    headers = {"Content-Disposition": f"attachment; filename={ascii_name}; filename*=UTF-8''{quote(filename)}"}
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


__all__ = ["router"]
